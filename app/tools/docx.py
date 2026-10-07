"""Microsoft Word document creation tool."""
from __future__ import annotations

from typing import Any
from pathlib import Path

from app.security.policy import RiskLevel
from app.security.sanitize import sanitize_tool_argument_string
from app.tools.base import Tool, ToolResult, ToolSpec
from app.tools.filesystem import _validate_path


class CreateDocxTool(Tool):
    spec = ToolSpec(
        name="create_docx",
        description="Create a Microsoft Word (.docx) document. LOW risk.",
        parameters_schema={
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Output path for .docx file"},
                "title": {"type": "string", "description": "Document title"},
                "content": {"type": "string", "description": "Document content (markdown-like or plain text)"},
            },
            "required": ["path", "title", "content"],
            "additionalProperties": False,
        },
        risk=RiskLevel.LOW,
    )

    def run(self, arguments: dict[str, Any]) -> ToolResult:
        try:
            path = _validate_path(arguments["path"])
            title = sanitize_tool_argument_string(arguments["title"], field_name="title")
            content = sanitize_tool_argument_string(arguments["content"], field_name="content")
        except Exception as e:
            return ToolResult(ok=False, payload=str(e))
        if not str(path).endswith('.docx'):
            path = Path(str(path) + '.docx')
        try:
            try:
                from docx import Document
            except ImportError:
                return ToolResult(ok=False, payload="python-docx not installed. Install with: pip install python-docx")
            doc = Document()
            doc.add_heading(title, 0)
            # split by lines and add paragraphs
            for line in content.split('\n'):
                line = line.strip()
                if not line:
                    continue
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('- '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                else:
                    doc.add_paragraph(line)
            doc.save(path)
            return ToolResult(ok=True, payload=f"Created {path}")
        except Exception as e:
            return ToolResult(ok=False, payload=str(e))
